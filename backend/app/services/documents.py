import io
import logging
import uuid
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.errors import AppError, NotFound
from app.db.models import Document, DocumentChunk, Project
from app.llm.gateway import get_gateway

log = logging.getLogger(__name__)

ALLOWED = {".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
           ".md": "text/markdown", ".txt": "text/plain"}
CHUNK_WORDS = 450  # ~600 tokens
OVERLAP_WORDS = 60


def extract_text(data: bytes, ext: str) -> str:
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in d.paragraphs)
    return data.decode("utf-8", errors="replace")


def chunk_text(text: str) -> list[str]:
    """Line-preserving chunking (~CHUNK_WORDS words) with ~OVERLAP_WORDS of trailing-line overlap."""
    lines = text.splitlines()
    chunks: list[str] = []
    cur: list[str] = []
    count = 0
    for line in lines:
        n = len(line.split())
        if cur and count + n > CHUNK_WORDS:
            chunks.append("\n".join(cur).strip())
            keep: list[str] = []
            kept = 0
            for prev in reversed(cur):
                kept += len(prev.split())
                if kept > OVERLAP_WORDS:
                    break
                keep.insert(0, prev)
            cur, count = keep, sum(len(x.split()) for x in keep)
        if n > CHUNK_WORDS:  # pathological single line: hard split
            words = line.split()
            for i in range(0, len(words), CHUNK_WORDS - OVERLAP_WORDS):
                chunks.append(" ".join(words[i:i + CHUNK_WORDS]))
            continue
        cur.append(line)
        count += n
    if "\n".join(cur).strip():
        chunks.append("\n".join(cur).strip())
    return [c for c in chunks if c]


def ingest(db: Session, project_id: uuid.UUID, filename: str, data: bytes, user_id: uuid.UUID | None) -> Document:
    s = get_settings()
    if db.get(Project, project_id) is None:
        raise NotFound("Project")
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED:
        raise AppError("unsupported_type", f"Allowed file types: {', '.join(ALLOWED)}", 415)
    if len(data) > s.max_upload_mb * 1024 * 1024:
        raise AppError("too_large", f"Max upload size is {s.max_upload_mb} MB", 413)
    try:
        text = extract_text(data, ext)
    except Exception as e:
        raise AppError("parse_failed", f"Could not parse {ext} file", 422) from e
    if not text.strip():
        raise AppError("empty_document", "No text could be extracted from the document", 422)
    doc_id = uuid.uuid4()
    path = s.upload_dir / str(project_id) / f"{doc_id}{ext}"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    doc = Document(id=doc_id, project_id=project_id, filename=Path(filename).name[:300], mime=ALLOWED[ext],
                   storage_path=str(path), extracted_text=text, uploaded_by=user_id)
    db.add(doc)
    add_chunks(db, doc, text)
    db.commit()
    log.info("document ingested", extra={"extra": {"document_id": str(doc_id), "chars": len(text)}})
    return doc


def add_chunks(db: Session, doc: Document, text: str) -> list[DocumentChunk]:
    pieces = chunk_text(text)
    vecs = get_gateway().embed(pieces)
    rows = [DocumentChunk(document_id=doc.id, idx=i, text=t, embedding=v) for i, (t, v) in enumerate(zip(pieces, vecs, strict=True))]
    db.add_all(rows)
    db.flush()
    return rows


def list_chunks(db: Session, document_id: uuid.UUID) -> list[DocumentChunk]:
    return list(db.scalars(select(DocumentChunk).where(DocumentChunk.document_id == document_id).order_by(DocumentChunk.idx)))
