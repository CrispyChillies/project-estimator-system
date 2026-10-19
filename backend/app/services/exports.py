from __future__ import annotations

import csv
import io
import uuid

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import DeliveryRole, Module, Task, TaskCategory
from app.services import estimates as es
from app.services import money
from app.services.calibration import mid
from app.services.summary import ensure_summary, with_user_costs

HEADER = ["Module", "Task", "Category", "Role", "Min h", "Max h", "Mid h", "Rate / h", "Cost (mid)", "Currency",
          "Rate used (VND per USD)", "Confidence", "Origin"]


def _ctx(db: Session, user, currency: str | None) -> tuple[str, float, dict]:
    """Display currency, workspace USD/VND rate and the user's effective VND rates."""
    from app.services import fx
    from app.services.rates import effective_rates

    cur = (currency or getattr(user, "display_currency", None) or "VND").upper()
    cur = cur if cur in money.CURRENCIES else "VND"
    return cur, fx.get_rate(db), effective_rates(db, user)


def _amount(vnd: float, cur: str, rate: float) -> float:
    return round(money.from_vnd(vnd, cur, rate), 2 if cur == "USD" else 0)


def _rows(db: Session, estimate_id: uuid.UUID, user=None, currency: str | None = None) -> list[list]:
    cur, usd_vnd, rates = _ctx(db, user, currency)
    q = (select(Task, Module, TaskCategory.name, DeliveryRole)
         .join(Module, Module.id == Task.module_id)
         .outerjoin(TaskCategory, TaskCategory.id == Task.category_id)
         .outerjoin(DeliveryRole, DeliveryRole.id == Task.delivery_role_id)
         .where(Task.estimate_id == estimate_id).order_by(Module.position, Task.position))
    out = []
    for t, m, cat, role in db.execute(q):
        if not es.is_effectively_included(t, m.removed):
            continue
        rate = rates.get(role.id, 0) if role else 0
        out.append([m.name, t.title, cat or "", role.name if role else "", t.min_hours, t.max_hours,
                    mid(t.min_hours, t.max_hours), _amount(rate, cur, usd_vnd),
                    _amount(mid(t.min_hours, t.max_hours) * rate, cur, usd_vnd), cur, usd_vnd, t.confidence,
                    t.origin])
    return out


def to_csv(db: Session, estimate_id: uuid.UUID, user=None, currency: str | None = None) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(HEADER)
    w.writerows(_rows(db, estimate_id, user, currency))
    return buf.getvalue()


def to_xlsx(db: Session, estimate_id: uuid.UUID, user=None, currency: str | None = None) -> bytes:
    est = es.get_estimate(db, estimate_id)
    cur, usd_vnd, _ = _ctx(db, user, currency)
    rows = _rows(db, estimate_id, user, currency)
    wb = Workbook()
    flat = wb.active
    flat.title = "Tasks"
    flat.append(HEADER)
    for c in flat[1]:
        c.font = Font(bold=True)
    for r in rows:
        flat.append(r)
    st = wb.create_sheet("Structured")
    current = None
    for r in rows:
        if r[0] != current:
            current = r[0]
            st.append([current])
            st.cell(st.max_row, 1).font = Font(bold=True)
        st.append(["", *r[1:]])
    sm = wb.create_sheet("Summary")
    s = with_user_costs(db, ensure_summary(db, est), user)
    sim = s.get("simulation", {})
    for k, v in [("Version", est.version), ("Status", est.status), ("Total min", s.get("total_min")),
                 ("Total max", s.get("total_max")), ("P50", sim.get("p50")), ("P80", sim.get("p80")),
                 ("P90", sim.get("p90")), ("Currency", cur), ("Rate used (VND per USD)", usd_vnd),
                 ("Cost (min)", _amount(s.get("total_cost_min", 0), cur, usd_vnd)),
                 ("Cost (mid)", _amount(s.get("total_cost_mid", 0), cur, usd_vnd)),
                 ("Cost (max)", _amount(s.get("total_cost_max", 0), cur, usd_vnd))]:
        sm.append([k, v])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
